import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Set background color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding for a table cell (values in dxa: 20 dxa = 1 pt)."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D7DE", sz="4", val="single"):
    """Set subtle modern borders for a table."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
                            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                            f'<w:insideV w:val="none"/>'
                            f'<w:left w:val="none"/>'
                            f'<w:right w:val="none"/>'
                            f'</w:tblBorders>')
        tblPr[0].append(borders)

def add_callout(doc, text, title=None, border_color="1E3A8A", bg_color="F0F4F8"):
    """Add a shaded callout block with a thick colored left border."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=130, bottom=130, left=180, right=180)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
                        f'<w:top w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'<w:bottom w:val="none"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if title:
        run_title = p.add_run(f"{title}\n")
        run_title.bold = True
        run_title.font.name = "Arial"
        run_title.font.size = Pt(10)
        r = int(border_color[0:2], 16)
        g = int(border_color[2:4], 16)
        b = int(border_color[4:6], 16)
        run_title.font.color.rgb = RGBColor(r, g, b)
    
    run_text = p.add_run(text)
    run_text.font.name = "Arial"
    run_text.font.size = Pt(9.5)
    run_text.font.italic = True
    run_text.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_box(doc, text):
    """Add a monospace diagram or architecture flow box."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:left w:val="single" w:sz="12" w:space="0" w:color="94A3B8"/>'
                        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_custom_table(table, col_widths, headers, data_rows, header_bg="1E3A8A"):
    """Format table with corporate styling, zebra striping, and cell margins."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="D0D7DE", sz="4")
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = Inches(col_widths[i])
        set_cell_background(hdr_cells[i], header_bg)
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(title)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    # Data Rows
    for r_idx, row_data in enumerate(data_rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].width = Inches(col_widths[c_idx])
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=70, bottom=70, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(30, 41, 59)

def add_embedded_image(doc, img_path, caption_text, width_inches=6.4):
    """Embed an image with a centered caption block."""
    if not os.path.exists(img_path):
        print(f"Warning: Image not found at {img_path}")
        return
        
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(10)
    p_img.paragraph_format.space_after = Pt(3)
    p_img.paragraph_format.keep_with_next = True
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(width_inches))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(0)
    p_cap.paragraph_format.space_after = Pt(12)
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = "Arial"
    run_cap.font.size = Pt(8.5)
    run_cap.font.italic = True
    run_cap.font.color.rgb = RGBColor(100, 116, 139)

def build_eda_framework_docx():
    current_dir = os.path.dirname(os.path.abspath(__file__))
    assets_dir = os.path.join(current_dir, "assets")
    output_docx = os.path.join(current_dir, "EDA_and_Visualization_Framework.docx")
    
    doc = docx.Document()
    
    # Page Setup: Letter, 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header & Footer
        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.text = "DATA SCIENCE INTERNSHIP | WEEK 2: EDA & VISUALIZATION FRAMEWORK DESIGN"
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_hdr.style.font.name = "Arial"
        p_hdr.style.font.size = Pt(8)
        p_hdr.style.font.color.rgb = RGBColor(148, 163, 184)
        
        footer = section.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.text = "Confidential - Professional Portfolio Deliverable | Author: Laxmi Gupta"
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_ftr.style.font.name = "Arial"
        p_ftr.style.font.size = Pt(8)
        p_ftr.style.font.color.rgb = RGBColor(148, 163, 184)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("Comprehensive Exploratory Data Analysis (EDA) and Visualization Framework Design")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0, 51, 102) # Deep Navy

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("An Industrial-Grade, Dataset-Agnostic Blueprint for Empirical Data Understanding, Statistical Hypothesis Testing, and Model Readiness")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Metadata Card
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    set_table_borders(meta_table, color="CBD5E1", sz="4")
    
    meta_data = [
        ("Author & Lead Data Scientist:", "Laxmi Gupta (Data Science Intern)"),
        ("Curriculum Milestone:", "Week 2: Exploratory Data Analysis & Statistical Visualization Design"),
        ("Target Analytical Ecosystem:", "Python 3.11+ | Pandas | NumPy | SciPy | Statsmodels | Matplotlib | Seaborn | Plotly | Missingno"),
        ("Time Budget Allocation:", "33.0 Focused Engineering Hours (Target: 30–35 Hours)"),
        ("Document Classification:", "Standard Operating Procedure (SOP) & Technical Reference Blueprint")
    ]
    
    for idx, (label, val) in enumerate(meta_data):
        c0 = meta_table.cell(idx, 0)
        c1 = meta_table.cell(idx, 1)
        c0.width = Inches(2.3)
        c1.width = Inches(4.2)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.name = "Arial"
        r0.font.size = Pt(8.5)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(30, 41, 59)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = "Arial"
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = RGBColor(51, 65, 85)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Executive Summary Callout
    add_callout(
        doc,
        "\"Exploratory data analysis can never be the whole story, but nothing else can serve as the foundation stone.\" "
        "— John W. Tukey. \n\n"
        "This framework establishes a dataset-agnostic, enterprise-grade architecture for conducting exhaustive exploratory data "
        "analysis and crafting strategic statistical visualizations. Designed to prevent data leakage, identify multivariate anomalies, "
        "verify algorithmic assumptions, and bridge raw ingestion with predictive modeling, this document codifies a 6-stage lifecycle, "
        "deterministic missingness/outlier decision trees, an exhaustive visualization taxonomy, and a 4-tier governance reporting standard.",
        title="EXECUTIVE ARCHITECTURAL MANDATE & PHILOSOPHY",
        border_color="003366",
        bg_color="F0F4F8"
    )

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0, 51, 102)

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 58, 138)

    def add_body(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(30, 41, 59)
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.1
        r_b = p.add_run(bold_prefix)
        r_b.font.name = "Arial"
        r_b.font.size = Pt(9)
        r_b.font.bold = True
        r_b.font.color.rgb = RGBColor(30, 41, 59)
        r_t = p.add_run(text)
        r_t.font.name = "Arial"
        r_t.font.size = Pt(9)
        r_t.font.color.rgb = RGBColor(51, 65, 85)

    # -------------------------------------------------------------
    # SECTION 1: INTRODUCTION & THEORETICAL FOUNDATIONS
    # -------------------------------------------------------------
    add_h1("1. Introduction & Theoretical Foundations of EDA")
    add_body(
        "In modern enterprise data science, Exploratory Data Analysis (EDA) represents the critical empirical bridge between raw data "
        "ingestion and predictive model engineering. While machine learning algorithms are frequently perceived as autonomous learners, "
        "their optimization algorithms operate strictly within the mathematical constraints and statistical distributions of the input space. "
        "Flawed distributions, undetected multi-collinear features, unhandled missingness mechanisms, and structural outliers directly distort "
        "loss functions, destabilize matrix inversions, and generate catastrophic generalization errors in production."
    )
    add_h2("1.1 The Exploratory vs. Confirmatory Paradigm")
    add_body(
        "Traditional statistical analysis prioritized Confirmatory Data Analysis (CDA), wherein researchers formulate rigid null and alternative "
        "hypotheses before gathering sample observations. While CDA remains vital for causal inference and A/B test evaluation, complex real-world "
        "enterprise datasets rarely conform to neat parametric assumptions. EDA, formalized by John W. Tukey in 1977, provides an inductive, "
        "data-driven methodology that prioritizes graphical discovery, residual diagnostics, and iterative parameter checking."
    )
    
    add_code_box(doc, 
"""[Raw Data Ingestion & Storage]
             │
             ▼
┌──────────────────────────────────────────────┐
│  Exploratory Data Analysis (Inductive Discovery) │ ◄── What does the data reveal?
│  - Structural Profiling & Anomaly Auditing   │     (Modality, Skewness, Nulls)
└──────────────────────┬───────────────────────┘
                       │
                       ▼
┌──────────────────────────────────────────────┐
│ Confirmatory Data Analysis (Deductive Tests) │ ◄── Is the observed pattern
│  - Welch's t-test, ANOVA, Chi-Square, VIF    │     statistically significant?
└──────────────────────┬───────────────────────┘
                       │
                       ▼
┌──────────────────────────────────────────────┐
│ Feature Engineering & Predictive ML Pipeline │ ◄── Supervised Learning on Clean,
│  - Certified Transformations & Ensembles     │     Mathematically Sound Data
└──────────────────────────────────────────────┘"""
    )
    
    add_h2("1.2 Strategic Objectives & Role in the CRISP-DM Framework")
    add_body(
        "Within the Cross-Industry Standard Process for Data Mining (CRISP-DM), this EDA framework directly executes Phase 2 (Data Understanding) "
        "while establishing the mathematical foundation for Phase 3 (Data Preparation). The four primary strategic imperatives of this framework are:"
    )
    add_bullet("1. Uncovering Latent Data Structures: ", "Identifying multi-modal distributions, non-linear relationships, and emergent natural customer or behavioral sub-clusters.")
    add_bullet("2. Exposing Data Degeneracies: ", "Isolating sensor failures, logging truncations, survivorship bias, structural zeroes, and high-cardinality leakages.")
    add_bullet("3. Validating Algorithmic Assumptions: ", "Evaluating multivariate normality, homoscedasticity, and independence required by linear classifiers, GLMs, and neural architectures.")
    add_bullet("4. Generating Transformation Blueprints: ", "Providing empirical justification for power transforms (Box-Cox/Yeo-Johnson), robust scalers, and target encoding strategies.")

    # -------------------------------------------------------------
    # SECTION 2: 6-STAGE EDA PIPELINE
    # -------------------------------------------------------------
    add_h1("2. Strategic 6-Stage EDA Execution Pipeline")
    add_body(
        "To guarantee reproducibility, rigor, and standardization across disparate business domains, this framework organizes exploratory "
        "analysis into six sequential yet iterative operational stages:"
    )

    img1_path = os.path.join(assets_dir, "eda_workflow_pipeline.png")
    add_embedded_image(doc, img1_path, "Figure 1: End-to-End 6-Stage EDA Execution Pipeline & Feedback Loop Architecture (300 DPI)", width_inches=6.5)

    add_body(
        "The six stages operate in tight synergy, feeding empirical findings back into earlier phases whenever unhandled anomalies or structural "
        "defects are identified:"
    )
    add_bullet("Stage 01: Ingestion & Schema Auditing: ", "Validates column data types, assesses memory usage, casts categorical downcasts, and standardizes ISO-8601 temporal timestamps.")
    add_bullet("Stage 02: Data Integrity & Cleaning: ", "Identifies missingness mechanisms (MCAR/MAR/MNAR), eliminates semantic duplicate records, and rectifies structural encoding corruptions.")
    add_bullet("Stage 03: Univariate Feature Profiling: ", "Computes non-parametric scale metrics (IQR, MAD), measures skewness/kurtosis, and identifies univariate outlier boundaries.")
    add_bullet("Stage 04: Bivariate & Cross-Feature Interactions: ", "Calculates parametric and non-parametric correlation matrices, contingency tables, and target-association significance tests.")
    add_bullet("Stage 05: Multivariate & Dimensionality Analysis: ", "Evaluates multicollinearity using Variance Inflation Factors (VIF), explores interaction terms, and applies PCA / UMAP embeddings.")
    add_bullet("Stage 06: Synthesis & Modeling Gate: ", "Assembles feature transformation recipes, creates data dictionaries, and issues formal mathematical readiness certification for ML training.")

    # -------------------------------------------------------------
    # SECTION 3: DATA TAXONOMY & CLASSIFICATION
    # -------------------------------------------------------------
    add_h1("3. Universal Data Taxonomy & Feature Classification")
    add_body(
        "A rigorous exploratory analysis begins with precise mathematical classification of feature types. Statistical tests and graphical "
        "archetypes must align strictly with the mathematical measurement scale of the target attribute:"
    )

    tax_headers = ["Data Class", "Measurement Scale", "Descriptive Metrics", "Visual Representations", "Recommended Python dtypes"]
    tax_rows = [
        ["Continuous Numerical", "Ratio / Interval", "Mean, Std, Median, IQR, Skewness, Kurtosis", "Histogram + KDE, Box, Violin, Q-Q Plot", "float64, float32"],
        ["Discrete Numerical", "Count (Non-negative)", "Median, Mode, Dispersion Index, Zero-Count %", "Bar Chart, Rootogram, Frequency Polygon", "int64, int32, UInt16"],
        ["Nominal Categorical", "Unordered Discrete", "Mode, Relative Frequencies, Shannon Entropy", "Bar Chart, Pareto Chart (80/20), Treemap", "category, string"],
        ["Ordinal Categorical", "Ranked Discrete", "Median, Percentiles, Rank-Order Frequencies", "Ordered Bar, Diverging Likert Chart", "CategoricalDtype(ordered=True)"],
        ["Temporal / DateTime", "Cyclic / Interval", "Min, Max, Span, Seasonality, Autocorrelation", "Longitudinal Line, Seasonal Subseries", "datetime64[ns], PeriodIndex"],
        ["Unstructured Text", "Lexical / Semantic", "Token Count, Char Length, Lexical Diversity", "Word Length Dist, N-Gram Bar, TF-IDF UMAP", "string, pyarrow.string"],
        ["Spatial / Geospatial", "Spherical Coordinates", "Centroid, Bounding Box, Moran's I Autocorr", "Choropleth Map, Hexbin Spatial Scatter", "geopandas.GeoSeries"]
    ]
    tax_table = doc.add_table(rows=len(tax_rows) + 1, cols=len(tax_headers))
    format_custom_table(tax_table, [1.3, 1.1, 1.5, 1.4, 1.2], tax_headers, tax_rows)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SECTION 4: STATISTICAL EXPLORATION TECHNIQUES
    # -------------------------------------------------------------
    add_h1("4. In-Depth Statistical Exploration Techniques")
    add_body(
        "Exploratory techniques are classified along dimensional horizons, expanding from isolated attribute inspections into "
        "complex multivariate manifold learning."
    )
    
    add_h2("4.1 Univariate Profiling & Distribution Diagnostics")
    add_body(
        "Univariate analysis examines individual features in isolation to establish distributional properties and detect corrupt records:"
    )
    add_bullet("Measures of Location & Scale: ", "Rather than relying solely on the arithmetic mean and standard deviation (which are easily distorted by extreme observations), the framework evaluates robust non-parametric metrics: the Median, Interquartile Range (IQR = Q3 - Q1), and Median Absolute Deviation (MAD = median(|Xi - median(X)|)). For continuous variables, a 5% Trimmed Mean is also computed to evaluate tail leverage.")
    add_bullet("Distribution Shape & Higher Moments: ", "Fisher-Pearson standardized skewness (S) is calculated to identify distributional asymmetry; |S| > 1.0 indicates severe skewness requiring non-linear transformation. Excess kurtosis (K) measures tail weight relative to a Gaussian distribution (K = 0); heavy leptokurtic tails (K > 0) indicate high vulnerability to extreme outliers.")
    add_bullet("Normality & Goodness-of-Fit Tests: ", "The Shapiro-Wilk test (for sample sizes N < 5,000) and the D'Agostino-Pearson K^2 test (for larger samples) quantify deviations from Gaussian behavior, supported graphically by Quantile-Quantile (Q-Q) probability plots.")
    add_bullet("Categorical Diversity & Shannon Entropy: ", "For categorical attributes, diversity is quantified using Shannon Information Entropy (H(X) = -sum(p(xi) * log2(p(xi)))). Low entropy signals severe class concentration in a single dominant category, whereas high entropy indicates uniform dispersion across levels.")

    add_h2("4.2 Bivariate Interactions & Hypothesis Testing")
    add_body(
        "Bivariate analysis investigates pairwise relationships to uncover predictive associations and detect redundant feature pairs:"
    )
    add_bullet("Continuous-to-Continuous Pairs: ", "Pearson's linear correlation coefficient (r) measures linear co-dependence under bivariate normality assumptions. For monotonic non-linear relationships or data containing extreme tails, Spearman's rank correlation (rho) and Kendall's tau (tau) are computed. Furthermore, the Maximal Information Coefficient (MIC) is calculated to capture complex non-linear curves (e.g., sinusoidal, parabolic, or multi-modal clusters).")
    add_bullet("Categorical-to-Categorical Pairs: ", "Two-way contingency cross-tabulations are subjected to Pearson's Chi-Square Test of Independence (chi^2). To quantify the strength of association independent of table dimensions, Cramér's V is calculated, yielding a normalized metric bounded between 0.0 (complete independence) and 1.0 (complete deterministic dependency).")
    add_bullet("Continuous-to-Categorical Pairs: ", "Differences in continuous feature distributions across categorical subgroups are evaluated using Welch's Two-Sample t-test (two categories) or One-Way Analysis of Variance (ANOVA F-test, three or more categories). When normality or equal variance assumptions are violated, non-parametric alternatives—the Mann-Whitney U test and Kruskal-Wallis H test—are executed.")

    add_h2("4.3 Multivariate Analysis & Dimensionality Reduction")
    add_body(
        "Multivariate analysis explores co-dependencies among three or more features simultaneously, identifying redundant feature manifolds "
        "and evaluating latent clustering:"
    )
    add_bullet("Multicollinearity & Variance Inflation Factor (VIF): ", "High correlation among independent variables inflates the standard errors of regression coefficients, causing numerical instability. The framework computes the VIF for each feature: VIF_j = 1 / (1 - R_j^2). Features exhibiting VIF > 5.0 are flagged for collinearity investigation, while VIF > 10.0 signals severe multicollinearity requiring feature elimination or dimensionality reduction.")
    add_bullet("Principal Component Analysis (PCA): ", "PCA projects high-dimensional continuous features onto orthogonal eigenvectors that maximize captured variance. Scree plots and cumulative explained variance curves identify the intrinsic dimensionality of the dataset.")
    add_bullet("Non-Linear Manifold Learning (t-SNE & UMAP): ", "For complex non-linear manifolds, Uniform Manifold Approximation and Projection (UMAP) and t-Distributed Stochastic Neighbor Embedding (t-SNE) compress high-dimensional feature spaces into 2D/3D visual embeddings, revealing natural behavioral clusters and latent outlier islands.")

    # -------------------------------------------------------------
    # SECTION 5: MISSING DATA & OUTLIER REMEDIATION
    # -------------------------------------------------------------
    add_h1("5. Missing Data Diagnostics & Outlier Remediation Framework")
    add_body(
        "Improper handling of missing values and extreme observations is a primary source of data leakage and production degradation. "
        "This framework establishes deterministic decision rules for anomaly diagnosis and treatment:"
    )

    img2_path = os.path.join(assets_dir, "missing_data_outlier_decision_tree.png")
    add_embedded_image(doc, img2_path, "Figure 2: Deterministic Decision Architecture for Missing Data Mechanisms & Outlier Treatment (300 DPI)", width_inches=6.5)

    add_h2("5.1 Missingness Mechanisms & Imputation Protocols")
    add_body(
        "Following Donald Rubin's foundational missingness taxonomy, missing values are classified into three distinct mathematical mechanisms:"
    )
    add_bullet("1. Missing Completely at Random (MCAR): ", "Missingness is completely independent of both observed features and unobserved values. Verified via Little's MCAR Test (p > 0.05). If missingness is below 5%, Listwise Deletion (Complete Case Analysis) produces statistically unbiased parameter estimates.")
    add_bullet("2. Missing at Random (MAR): ", "Missingness depends systematically on other observed features in the dataset, but not on the missing value itself (e.g., older respondents omitting income, while age is observed). Resolved via Multiple Imputation by Chained Equations (MICE), K-Nearest Neighbors (KNN) Imputation, or MissForest.")
    add_bullet("3. Missing Not at Random (MNAR): ", "Missingness depends directly on the unobserved missing value itself (e.g., high-income earners refusing to disclose wealth). Resolved by adding an explicit binary missingness indicator column (is_missing_feature_X), followed by median/mode imputation or pattern-mixture modeling.")

    add_callout(
        doc,
        "CRITICAL LEAKAGE PREVENTATIVE RULE: In all machine learning pipelines, imputation parameters (e.g., feature medians, "
        "modes, KNN distance trees, or MICE regression estimators) must be fitted EXCLUSIVELY on the training partition. "
        "Imputers are then applied as frozen transformers to validation and test folds to prevent data leakage.",
        title="GOVERNANCE RULE: IMPUTATION LEAKAGE GUARD",
        border_color="D97706",
        bg_color="FFFBEB"
    )

    add_h2("5.2 Outlier Detection & Treatment Hierarchy")
    add_body(
        "Outliers are systematically detected, classified, and treated based on their underlying generative mechanism:"
    )
    add_bullet("Tukey's IQR Outer & Inner Fences: ", "Inner fences [Q1 - 1.5*IQR, Q3 + 1.5*IQR] identify mild outliers; outer fences [Q1 - 3.0*IQR, Q3 + 3.0*IQR] flag extreme outliers. Completely distribution-free and non-parametric.")
    add_bullet("Modified Z-Score via MAD: ", "For moderately skewed distributions, the Modified Z-score (M_i = 0.6745 * (X_i - Median) / MAD) provides a resilient outlier metric. Values with |M_i| > 3.5 are flagged as statistical outliers.")
    add_bullet("Mahalanobis Distance: ", "Detects multivariate outliers by evaluating the distance of an observation from the multidimensional centroid, incorporating the covariance matrix: D_M(x) = sqrt((x - mu)^T * Sigma^(-1) * (x - mu)). Observations exceeding the 97.5th percentile of the Chi-Square distribution are marked as multivariate anomalies.")
    add_bullet("Isolation Forest & Local Outlier Factor (LOF): ", "For high-dimensional feature spaces, Isolation Forests isolate anomalies through random hyperplane splits, assigning high anomaly scores to observations with short average path lengths.")
    add_bullet("Remediation Hierarchy: ", "1) Measurement / System Errors: Discard or impute. 2) Valid Extreme Events: Apply Winsorization (capping at 1st and 99th percentiles) or non-linear power transformations (Log1p, Yeo-Johnson). 3) Modeling Strategy: Prioritize tree ensembles (XGBoost, LightGBM) that split on feature ranks rather than absolute magnitudes.")

    # -------------------------------------------------------------
    # SECTION 6: VISUALIZATION STRATEGY TAXONOMY
    # -------------------------------------------------------------
    add_h1("6. Strategic Visualization Taxonomy & Plotting Matrix")
    add_body(
        "Effective data visualization requires selecting visual encodings (position, length, area, angle, color) that optimize human visual "
        "perception while maximizing the data-ink ratio (Edward Tufte). The framework maps visual archetypes directly to analytical objectives:"
    )

    img3_path = os.path.join(assets_dir, "visualization_strategy_taxonomy.png")
    add_embedded_image(doc, img3_path, "Figure 3: Strategic Visualization Taxonomy & Statistical Plotting Matrix (300 DPI)", width_inches=6.5)

    add_body(
        "The following matrix formalizes visual selections across core analytical objectives, feature dimensionalities, and library implementations:"
    )

    viz_headers = ["Analytical Objective", "Feature Types", "Primary Visualization", "Statistical Insight & Business Utility", "Python API"]
    viz_rows = [
        ["Distribution & Spread", "1 Continuous", "Histogram + KDE Overlay", "Reveals modality (unimodal/bimodal), tail heaviness, skewness, and zero-inflation.", "seaborn.histplot(kde=True)"],
        ["Distribution Comparison", "1 Cont. + 1 Cat.", "Violin Plot + Box Inset", "Contrasts median, IQR, whiskers, and full kernel density distribution across groups.", "seaborn.violinplot()"],
        ["Normality Verification", "1 Continuous", "Normal Q-Q Probability Plot", "Directly compares empirical quantiles against theoretical standard normal quantiles.", "statsmodels.api.qqplot(line='s')"],
        ["Cumulative Density", "1 Continuous", "Empirical CDF (ECDF)", "Visualizes exact percentile values without binning artifacts or bandwidth distortions.", "seaborn.ecdfplot()"],
        ["Category Benchmarking", "1 Categorical", "Pareto Chart (80/20 Rule)", "Combines descending frequency bars with cumulative percentage curve to spot key drivers.", "matplotlib dual-axis"],
        ["Hierarchical Shares", "2+ Hierarchical Cats", "Treemap & Sunburst", "Represents nested category proportions using area-based geometric partitioning.", "plotly.express.treemap()"],
        ["Bivariate Dependency", "2 Continuous", "Joint Scatter + Marginal KDE", "Plots continuous co-variance alongside univariate boundary density distributions.", "seaborn.jointplot(kind='reg')"],
        ["Dense Bivariate ($N > 10^5$)", "2 Continuous", "Hexagonal Binning (Hexbin)", "Mitigates scatter plot overplotting through 2D spatial binning and color density mapping.", "matplotlib.pyplot.hexbin()"],
        ["Correlation Matrix", "$N$ Continuous", "Clustered Heatmap", "Reorders feature correlation matrix using hierarchical clustering to reveal feature collinearity.", "seaborn.clustermap()"],
        ["Multi-Feature Interaction", "3–6 Mixed Features", "Facet Grid Pairplot", "Renders an N x N matrix of scatter and density distributions conditioned on target hue.", "seaborn.pairplot(hue=...)"],
        ["Temporal Decomposition", "1 DateTime + 1 Cont.", "Classical Time-Series Decomp", "Separates longitudinal signals into Trend, Seasonality, and Cyclical/Residual noise.", "statsmodels.tsa.seasonal_decompose"],
        ["High-Dim Manifold", "4+ Continuous", "Parallel Coordinates Plot", "Traces multidimensional entity paths across parallel axes to identify cluster profiles.", "plotly.express.parallel_coordinates()"]
    ]
    viz_table = doc.add_table(rows=len(viz_rows) + 1, cols=len(viz_headers))
    format_custom_table(viz_table, [1.3, 1.1, 1.4, 1.6, 1.1], viz_headers, viz_rows)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SECTION 7: PYTHON TOOLING & ARCHITECTURE
    # -------------------------------------------------------------
    add_h1("7. Python Tooling, Library Ecosystem & Architecture")
    add_body(
        "The Python open-source data science stack provides specialized, high-performance libraries across each layer of the exploratory lifecycle. "
        "The architecture divides responsibilities across dedicated functional engines:"
    )

    py_tools = [
        ("Pandas & NumPy: ", "The foundational vectorized data manipulation layer. Used for memory-efficient downcasting, grouped aggregations, windowing operations, missingness masks, and structural assertions."),
        ("SciPy & Statsmodels: ", "The rigorous inferential statistical layer. Executes parametric and non-parametric hypothesis tests (Shapiro-Wilk, Mann-Whitney U, Kruskal-Wallis), computes VIF scores, and performs time-series decomposition."),
        ("Matplotlib & Seaborn: ", "The primary static publication graphics layer. Matplotlib provides low-level canvas layout, axis orchestration, and DPI control; Seaborn delivers high-level statistical visualizations with built-in confidence intervals and elegant themes."),
        ("Plotly Express & Dash: ", "The dynamic interactive visualization layer. Renders GPU-accelerated WebGL charts, interactive geographic choropleths, and multi-dimensional parallel coordinate plots for stakeholder discovery."),
        ("Missingno: ", "Dedicated nullity visualization tool. Renders missingness matrices, correlation heatmaps, and hierarchical nullity dendrograms to quickly diagnose missingness mechanisms."),
        ("YData-Profiling & Sweetviz: ", "Automated exploratory profiling engines. Generates comprehensive HTML baseline dossiers detailing feature types, zero-variance attributes, correlations, and warnings in seconds.")
    ]
    for p_name, p_desc in py_tools:
        add_bullet(p_name, p_desc)

    add_body(
        "Below are production-ready Python implementations for automated schema auditing and Variance Inflation Factor (VIF) evaluation:"
    )

    add_code_box(doc,
"""import pandas as pd
import numpy as np
from statsmodels.stats.outliers_influence import variance_inflation_factor

# 1. Automated Schema & Memory Profiling Engine
def audit_dataset_health(df: pd.DataFrame) -> pd.DataFrame:
    \"\"\"Audits nullity, unique cardinality, memory footprint, and datatypes.\"\"\"
    audit = pd.DataFrame({
        'Dtype': df.dtypes,
        'Null_Count': df.isnull().sum(),
        'Null_Pct': np.round((df.isnull().sum() / len(df)) * 100, 2),
        'Unique_Values': df.nunique(),
        'Memory_MB': np.round(df.memory_usage(deep=True) / (1024 ** 2), 3)
    })
    return audit.sort_values(by='Null_Pct', ascending=False)

# 2. Multicollinearity VIF Diagnostic Engine
def compute_multicollinearity_vif(X: pd.DataFrame, threshold: float = 5.0) -> pd.DataFrame:
    \"\"\"Calculates VIF across numerical predictors and flags collinear features.\"\"\"
    X_clean = X.dropna()
    vif_df = pd.DataFrame()
    vif_df["Feature"] = X_clean.columns
    vif_df["VIF"] = [variance_inflation_factor(X_clean.values, i) for i in range(X_clean.shape[1])]
    vif_df["Collinear_Flag"] = vif_df["VIF"] > threshold
    return vif_df.sort_values(by="VIF", ascending=False)"""
    )

    # -------------------------------------------------------------
    # SECTION 8: REPORTING & GOVERNANCE ARCHITECTURE
    # -------------------------------------------------------------
    add_h1("8. Reporting, Governance & Stakeholder Communication Architecture")
    add_body(
        "An exploratory analysis achieves maximum organizational value only when its findings are effectively communicated to both technical "
        "implementers and non-technical executive stakeholders. This framework enforces a 4-tier reporting hierarchy:"
    )

    img4_path = os.path.join(assets_dir, "eda_reporting_governance_architecture.png")
    add_embedded_image(doc, img4_path, "Figure 4: 4-Tier EDA Artifact Hierarchy Bridging Technical Rigor and Strategic Decision-Making (300 DPI)", width_inches=6.5)

    add_body(
        "The four communication tiers serve distinct stakeholder groups across the project lifecycle:"
    )
    add_bullet("Tier 1: Automated Baseline Health Dossier: ", "Generated within minutes of ingestion via YData-Profiling or Sweetviz. Provides data engineers and scientists with rapid feedback on schema health, duplicated keys, zero-variance columns, and missingness density.")
    add_bullet("Tier 2: Technical Deep-Dive Jupyter Blueprint: ", "A fully reproducible, version-controlled notebook (.ipynb) containing mathematical proofs, statistical test statistics (p-values, F-scores), VIF pruning iterations, and code-level transformation recipes. Serves as the technical specification for peer review and ML pipeline engineering.")
    add_bullet("Tier 3: Interactive Exploratory Web Application: ", "A Streamlit or Plotly Dash web dashboard enabling business analysts and domain experts to dynamically slice and dice cohorts, adjust what-if sliders, and validate behavioral patterns without writing code.")
    add_bullet("Tier 4: Executive Strategy Dossier & Modeling Gate: ", "A formal executive Word/PDF report summarizing primary business takeaways, data quality scorecards, risk disclosures, and the official sign-off authorizing the project to advance to predictive model training.")

    # -------------------------------------------------------------
    # SECTION 9: 33-HOUR TIMELINE ALLOCATION
    # -------------------------------------------------------------
    add_h1("9. Comprehensive 33-Hour Timeline & Task Breakdown")
    add_body(
        "To ensure thorough execution and maintain rigorous technical discipline, the exploratory framework is budgeted across "
        "33.0 focused engineering hours (perfectly matching the 30–35 hour requirement):"
    )

    sched_headers = ["Module & Milestone", "Core Analytical Activities", "Hours Allocated", "Cumulative %", "Key Deliverable Artifact"]
    sched_rows = [
        ["M1: Ingestion & Schema Auditing", "Data ingestion, memory profiling, type casting, datetime parsing, unique key validation.", "5.0h", "15.2%", "Automated Ingestion Profile & Schema Data Dictionary"],
        ["M2: Missing Data & Outliers", "Mechanism diagnosis (Little's test, missingno), imputation benchmarking, outlier detection (IQR, MAD, Mahalanobis).", "6.5h", "34.8%", "Anomaly Resolution Protocol & Imputation Pipeline Code"],
        ["M3: Univariate Distribution Profiling", "Central tendency, dispersion, skewness/kurtosis calculations, normality tests, Pareto charts, ECDF curves.", "6.0h", "53.0%", "Univariate Feature Profiles & Transformation Recommendations"],
        ["M4: Bivariate & Hypothesis Testing", "Pearson/Spearman matrices, Chi-Square, Welch's t-tests, ANOVA, target correlation ranking, scatter/density grids.", "6.5h", "72.7%", "Bivariate Statistical Matrix & Initial Feature Hypotheses"],
        ["M5: Multivariate & Dimensionality Analysis", "Multicollinearity VIF screening, PCA decomposition, t-SNE/UMAP projections, unsupervised cluster identification.", "5.0h", "87.9%", "Dimension Reduction Maps & VIF Pruning Blueprint"],
        ["M6: Synthesis, Visualization & Reporting", "Compilation of 4-tier reporting artifacts, executive deliverable authoring, modeling sign-off review.", "4.0h", "100.0%", "Final Executive DOCX Deliverable & Interactive Dashboard"],
        ["TOTAL ALLOCATION", "Comprehensive End-to-End Enterprise EDA Lifecycle", "33.0h", "100.0%", "Production-Ready ML Modeling Gate Certification"]
    ]
    sched_table = doc.add_table(rows=len(sched_rows) + 1, cols=len(sched_headers))
    format_custom_table(sched_table, [1.4, 1.8, 0.9, 0.9, 1.5], sched_headers, sched_rows)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SECTION 10: ANTICIPATED CHALLENGES & MITIGATIONS
    # -------------------------------------------------------------
    add_h1("10. Anticipated Challenges & Strategic Mitigations")
    add_body(
        "Enterprise exploratory data analysis routinely encounters complex operational hurdles. The framework incorporates proactive mitigations:"
    )

    risk_headers = ["Operational Challenge", "Latent Analytical Risk", "Architectural Mitigation Strategy"]
    risk_rows = [
        ["Big Data Scalability ($N > 10^7$)", "RAM exhaustion, slow plotting, and kernel crashes.", "Utilize chunked out-of-core engines (Polars, PyArrow), downcast numeric dtypes, and deploy 2D Hexbins or Datashader over raw scatter plots."],
        ["Data Leakage into Test Folds", "Artificially inflated validation scores; collapse in production.", "Enforce strict scikit-learn Pipeline encapsulation where imputers, scalers, and encoders fit exclusively on training folds."],
        ["High-Cardinality Categoricals", "Exploding sparse matrices from One-Hot Encoding; high memory.", "Consolidate rare levels into an 'Other' bin; utilize Target Encoding with empirical Bayes shrinkage and out-of-fold regularization."],
        ["Spurious Multiple Correlations", "Inflated Type I error rates when evaluating hundreds of pairwise features.", "Apply Bonferroni corrections or Benjamini-Hochberg False Discovery Rate (FDR) thresholds during batch hypothesis testing."],
        ["Extreme Class Imbalance (< 1%)", "Standard correlation and scatter plots conceal minority patterns.", "Evaluate Precision-Recall curves, apply stratified sampling during visual profiling, and inspect class-conditioned probability density functions."]
    ]
    risk_table = doc.add_table(rows=len(risk_rows) + 1, cols=len(risk_headers))
    format_custom_table(risk_table, [1.6, 1.8, 3.1], risk_headers, risk_rows)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SECTION 11: CONCLUSION & NEXT STEPS
    # -------------------------------------------------------------
    add_h1("11. Conclusion & Transition to Feature Engineering")
    add_body(
        "This Exploratory Data Analysis and Visualization Framework provides a standardized, mathematically rigorous foundation for extracting "
        "actionable insights and preparing datasets for enterprise predictive modeling. By transitioning systematically from automated baseline "
        "reconnaissance through univariate, bivariate, and multivariate analysis to a formal 4-tier reporting architecture, data teams eliminate "
        "data leakage, avoid costly model retraining cycles, and ensure transparent stakeholder communication."
    )
    add_body(
        "In Week 3, this framework will directly interface with downstream Feature Engineering and Data Preprocessing Pipelines, translating "
        "the empirical transformation recommendations established in this dossier into production-grade Scikit-Learn transformers."
    )

    # Save document
    doc.save(output_docx)
    print(f"Successfully generated DOCX deliverable at: {output_docx}")

if __name__ == "__main__":
    build_eda_framework_docx()
